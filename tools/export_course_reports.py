"""Export current course research and SPEC Markdown sources to Word."""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
FILES = [
    ROOT / "docs/research/学术成果分享平台需求调研报告V1.0.md",
    ROOT / "docs/research/调研依据与执行记录V1.0.md",
    ROOT / "docs/spec/学术成果分享平台核心规则SPEC初稿V0.1.md",
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        item = borders.find(qn(tag))
        if item is None:
            item = OxmlElement(tag)
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "3")
        item.set(qn("w:color"), "D7DCE2")


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, value, size=None):
    # Preserve readable labels and addresses instead of exposing Markdown syntax.
    value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", value)
    for part in re.split(r"(`[^`]+`)", value):
        if not part:
            continue
        is_code = part.startswith("`") and part.endswith("`")
        run = paragraph.add_run(part[1:-1] if is_code else part)
        set_run_font(run, size=size, bold=is_code)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(32, 38, 44)
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.space_after = Pt(6)

    for level, pt, before, after in [(1, 15, 16, 7), (2, 12, 11, 5)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(pt)
        style.font.bold = True
        style.font.color.rgb = RGBColor(20, 28, 36)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(18, 25, 31)
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("OvO 小组  ·  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        set_run_font(run, 8, color="777F86")


def add_table(doc, rows):
    cols = len(rows[0])
    table = doc.add_table(rows=1, cols=cols)
    table.autofit = False
    table.style = "Table Grid"
    total_width = Cm(16.6)
    widths = [total_width / cols] * cols
    if cols == 2:
        widths = [Cm(4.2), Cm(12.4)]
    elif cols == 3:
        widths = [Cm(3.4), Cm(6.6), Cm(6.6)]
    elif cols == 4:
        widths = [Cm(2.65), Cm(4.65), Cm(4.65), Cm(4.65)]
    for i, row in enumerate(rows):
        cells = table.rows[0].cells if i == 0 else table.add_row().cells
        for j, content in enumerate(row):
            cell = cells[j]
            cell.width = widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if i == 0:
                set_cell_shading(cell, "E9EDF1")
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.space_before = Pt(1)
            para.paragraph_format.line_spacing = 1.13
            if i == 0:
                para.paragraph_format.keep_with_next = True
            add_inline(para, content, 8.4 if cols >= 3 else 9)
            if i == 0:
                for run in para.runs:
                    run.bold = True
        tr_pr = table.rows[i]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
        if i == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].startswith("|"):
        cells = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
            rows.append(cells)
        i += 1
    return rows, i


def export(source):
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure(doc)
    paragraph_buffer = []

    def flush():
        if paragraph_buffer:
            para = doc.add_paragraph()
            add_inline(para, " ".join(paragraph_buffer))
            paragraph_buffer.clear()

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            flush()
            i += 1
            continue
        if line.startswith("|"):
            flush()
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if line.startswith("# "):
            flush()
            para = doc.add_paragraph(style="Title")
            add_inline(para, line[2:])
        elif line.startswith("## "):
            flush()
            para = doc.add_paragraph(style="Heading 1")
            add_inline(para, line[3:])
            if source.name.startswith("学术成果分享平台需求调研报告") and line.startswith("## 8."):
                para.paragraph_format.page_break_before = True
        elif line.startswith("### "):
            flush()
            para = doc.add_paragraph(style="Heading 2")
            add_inline(para, line[4:])
        elif re.match(r"^\d+\. ", line):
            flush()
            para = doc.add_paragraph(style="List Number")
            add_inline(para, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            flush()
            para = doc.add_paragraph(style="List Bullet")
            add_inline(para, line[2:])
        else:
            paragraph_buffer.append(line)
        i += 1
    flush()
    destination = source.with_suffix(".docx")
    doc.save(destination)
    print(destination)


if __name__ == "__main__":
    for source in FILES:
        export(source)
