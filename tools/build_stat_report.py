from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[1]
OUT = REPO_ROOT / "docs" / "research" / "学术成果分享平台现成调研统计报告.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

BLUE = "17365D"
LIGHT_BLUE = "EAF1F8"
PALE_BLUE = "F5F8FC"
GRAY = "D9D9D9"
TEXT = RGBColor(35, 42, 52)
MUTED = RGBColor(92, 104, 119)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=GRAY, size="5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        p_bdr.append(node)
    p_pr.append(p_bdr)


def set_font(run, east="Microsoft YaHei", latin="Aptos", size=10.5, bold=False, color=TEXT):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5A93")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for node in (fld_begin, instr, fld_sep, text, fld_end):
        run._r.append(node)
    set_font(run, size=9, color=MUTED)


def add_paragraph(doc, text="", bold_lead=None, style=None, align=None, space_after=5, line=1.35):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def set_table_widths(table, widths_cm):
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def style_table(table, widths, header_fill=BLUE, font_size=8.8, alignments=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_widths(table, widths)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        prevent_row_split(row)
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            if r_idx == 0:
                set_cell_shading(cell, header_fill)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                if alignments:
                    p.alignment = alignments[c_idx]
                for run in p.runs:
                    set_font(run, size=font_size, bold=(r_idx == 0), color=(RGBColor(255, 255, 255) if r_idx == 0 else TEXT))


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.75)
sec.bottom_margin = Cm(1.65)
sec.left_margin = Cm(1.8)
sec.right_margin = Cm(1.8)
sec.header_distance = Cm(0.7)
sec.footer_distance = Cm(0.7)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.35

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
title_style.font.size = Pt(25)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(6)

for name, size, before, after in (("Heading 1", 16, 12, 6), ("Heading 2", 12, 8, 4)):
    st = styles[name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run("软件系统分析与设计  需求调研支撑材料")
set_font(hr, size=8.5, color=MUTED)
add_page_number(sec.footer.paragraphs[0])

doc.core_properties.title = "学术成果分享平台现成调研统计报告"
doc.core_properties.subject = "需求调研统计汇总"
doc.core_properties.author = ""

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("学术成果分享平台\n现成调研统计报告")
remove_paragraph_borders(p)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)
r = p.add_run("课程实践需求调研支撑材料  |  2026年10月修订")
set_font(r, size=10.5, color=MUTED)

add_paragraph(
    doc,
    "本报告汇总可追溯的公开二手研究，用于辅助识别学术成果分享平台的需求。各研究的抽样对象、时间与统计口径不同，不能相加为统一样本，也不能据此推断本校用户的实际比例。当前较有依据的设计关注点是成果发现与展示、维护负担和指标可解释性；具体优先级仍须与课程任务边界和本组需求文档共同判断。",
    space_after=8,
)

doc.add_heading("核心结论", level=1)
for lead, rest in [
    ("成果发现优先。", "科研人员主要使用平台查找、分享和展示成果，深度问答与评论的使用率较低。"),
    ("维护成本必须低。", "在一项ResearchGate调查中，83%的受访者每周检索或浏览该平台不足1小时，94%每周维护个人主页不足1小时；这不是对全部平台的统一估计。"),
    ("可信性高于功能数量。", "作者消歧、成果归属、来源追踪和统计口径透明直接影响平台可信度。"),
    ("评价指标必须可解释。", "同一调查中，57%的受访者表示不理解ResearchGate Score的计算方式；本项目应明确展示统计口径。"),
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r1 = p.add_run(lead)
    set_font(r1, bold=True)
    r2 = p.add_run(rest)
    set_font(r2)

doc.add_heading("一  跨研究统计结果", level=1)

rows = [
    ("Nature全球调查", ">3,500人，95国", "认知与访问", "88%听说过ResearchGate；接近50%定期访问", "存在稳定的平台认知和检索需求"),
    ("PLOS ONE", "695份有效问卷", "使用时间", "83%每周检索/浏览RG不足1小时；94%每周主页维护不足1小时", "减少维护步骤，但比例不外推"),
    ("奥地利科研调查", "379人", "时间障碍", "47%认为时间不足阻碍进一步使用", "降低操作和维护负担"),
    ("奥地利科研调查", "ResearchGate用户", "科研可见度", "66%认为平台提高了自身可见度", "个人门户和成果展示有明确价值"),
    ("PLOS ONE", "695名科研人员", "指标透明度", "57%不理解RG Score的计算方式", "避免黑箱综合评分"),
    ("中文Kano研究", "183份有效问卷", "需求属性", "文献获取Better=0.748、Worse=-0.614", "检索和获取体验值得优先验证"),
]
table = doc.add_table(rows=1, cols=5)
headers = ["研究", "样本", "统计维度", "主要结果", "需求含义"]
for i, text in enumerate(headers):
    table.rows[0].cells[i].text = text
for row in rows:
    cells = table.add_row().cells
    for i, text in enumerate(row):
        cells[i].text = text
style_table(
    table,
    [2.7, 2.7, 2.4, 4.5, 4.2],
    font_size=8.2,
    alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
)

add_paragraph(doc, "注：RG为ResearchGate。不同研究的抽样地区、年份、对象和问题不同；不合并样本量，不机械平均比例，也不把研究结论等同于本组一手调研。", space_after=3)

doc.add_heading("二  中国用户需求优先级", level=1)
add_paragraph(
    doc,
    "2024年中文研究先访谈12名科研用户，再回收221份问卷并获得183份有效问卷，有效率82.8%。研究使用Kano模型和Better-Worse系数识别需求属性。",
    space_after=6,
)

kano = [
    ("项目申报", "0.785", "-0.591", "魅力型", "创新增强"),
    ("激励机制", "0.782", "-0.508", "魅力型", "创新增强"),
    ("多端使用", "0.756", "-0.717", "期望型", "体验优化"),
    ("文献获取", "0.748", "-0.614", "期望型", "核心功能"),
    ("社交网络", "0.667", "-0.727", "期望型", "合作发现"),
    ("科研协同", "0.661", "-0.761", "期望型", "合作发现"),
    ("交流互动", "0.453", "-0.807", "必备型", "基础能力"),
    ("学术评价", "0.279", "-0.745", "必备型", "透明统计"),
    ("自我展示", "0.491", "-0.304", "无差别型", "控制投入"),
    ("写作投稿", "0.367", "-0.391", "无差别型", "后续考虑"),
]
table2 = doc.add_table(rows=1, cols=5)
for i, text in enumerate(["需求", "Better", "Worse", "类型", "平台判断"]):
    table2.rows[0].cells[i].text = text
for row in kano:
    cells = table2.add_row().cells
    for i, text in enumerate(row):
        cells[i].text = text
style_table(
    table2,
    [4.0, 2.5, 2.5, 3.2, 4.3],
    font_size=9,
    alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
)

add_paragraph(doc, "Better越高，提供该功能越能提高满意度；Worse绝对值越高，缺少该功能越容易引发不满。", space_after=4)

h3 = doc.add_heading("三  证据解释边界", level=1)
h3.paragraph_format.page_break_before = True
add_paragraph(
    doc,
    "上述数据分别来自ResearchGate用户、一般科研人员及中文学术社交网络需求研究。它们能够提示潜在痛点，却不能直接给出本校师生的占比或证明本项目某功能必然受欢迎。Kano的Better/Worse系数也不等同于使用人数比例。",
    space_after=6,
)

doc.add_heading("四  综合需求判断", level=1)

judgments = [
    ("一期核心", "统一检索、多源成果聚合、作者消歧、科研人员门户、成果认领、来源与更新时间展示、透明统计。"),
    ("体验重点", "自动生成候选成果列表，支持批量确认、批量排除和异常纠错，减少科研人员长期维护成本。"),
    ("评价原则", "展示成果数、引用数、年度趋势和合作关系等可解释指标，不生成无法复核的统一综合分。"),
    ("学科适配", "统计比较应提示来源覆盖差异；一期成果类型以本组SRS规定的论文、专利、项目和软件著作为边界。"),
    ("后续增强", "合作发现和推荐属于需求验证方向；是否进入一期，以课程任务书与需求基线为准。"),
    ("控制范围", "保留轻量联系与反馈能力，但一期不建设完整的科研社交社区、复杂积分体系或写作投稿工具。"),
]
for title, text in judgments:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(title + "  ")
    set_font(r1, bold=True)
    r2 = p.add_run(text)
    set_font(r2)

doc.add_heading("五  统计口径说明", level=1)
add_paragraph(
    doc,
    "本报告是公开二手资料汇编，未实际执行本组访谈或问卷。不同研究的统计不能加总、平均或作为本校人群估计。详细检索、来源与执行状态见《调研依据与执行记录V1.0》。使用这些结果时，先核对原始问题和样本，再由本组需求边界筛选；本报告中的设计建议并非额外已批准需求。",
    space_after=5,
)

ref_heading = doc.add_heading("参考资料", level=1)
ref_heading.paragraph_format.page_break_before = True
refs = [
    ("Van Noorden R. Online collaboration: Scientists and the social network. Nature, 2014.", "https://www.nature.com/news/online-collaboration-scientists-and-the-social-network-1.15711"),
    ("Meier A, Tunger D. Survey on opinions and usage patterns for the ResearchGate platform. PLOS ONE, 2018.", "https://doi.org/10.1371/journal.pone.0204945"),
    ("吴瑞, 张耀坤. 基于用户信息需求分析的学术社交网络功能优化策略研究. 数字图书馆论坛, 2024.", "https://dlf.istic.ac.cn/dlf/ch/reader/view_abstract.aspx?file_no=202408006&flag=1"),
    ("Buehling K et al. Competitive performativity of academic social networks. Research Evaluation, 2024.", "https://doi.org/10.1093/reseval/rvae048"),
]
for idx, (citation, url) in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{idx}. {citation} ")
    set_font(r, size=8.7)
    add_hyperlink(p, "链接", url)

doc.save(OUT)
print(OUT)
