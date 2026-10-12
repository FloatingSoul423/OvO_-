"""Export the reviewed Markdown SRS bundle to one course-submission Word file.

Markdown remains the canonical source. Run after build_srs_review.py and
check_srs_quality.py; this script does not infer implementation or diagram status.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SRS = ROOT / "docs" / "srs"
OUTPUT = SRS / "OvO-SRS-V0.4.docx"


def clean(value: str) -> str:
    value = re.sub(r"<a\s+id=[^>]+></a>", "", value)
    value = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1", value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"\1", value)
    value = value.replace("`", "").replace("<br>", "；")
    return value.strip()


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    item = OxmlElement("w:cantSplit")
    tr_pr.append(item)


def add_table(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        parts = [clean(c) for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", p or "") for p in parts):
            continue
        rows.append(parts)
    if not rows:
        return
    count = max(map(len, rows))
    table = doc.add_table(rows=0, cols=count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for ridx, cells in enumerate(rows):
        row = table.add_row()
        no_split(row)
        for idx in range(count):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = cells[idx] if idx < len(cells) else ""
            if ridx == 0:
                shade(cell, "E6EDF5")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            elif ridx % 2 == 0:
                shade(cell, "F7F9FC")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_markdown(doc: Document, path: Path, *, base: int = 0, skip_title: bool = False) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    table_lines: list[str] = []
    first_heading = True
    for line in lines + [""]:
        if line.startswith("|"):
            table_lines.append(line)
            continue
        if table_lines:
            add_table(doc, table_lines)
            table_lines = []
        if not line.strip() or line.startswith("<a id="):
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1)) + base
            if skip_title and first_heading:
                first_heading = False
                continue
            first_heading = False
            label = clean(m.group(2))
            if level <= 3:
                p = doc.add_paragraph(label, style=f"Heading {max(1, level)}")
                p.paragraph_format.keep_with_next = True
            else:
                p = doc.add_paragraph(style="Normal")
                run = p.add_run(label)
                run.bold = True
                run.font.color.rgb = RGBColor(31, 78, 121)
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(7)
            continue
        if line.startswith(">"):
            p = doc.add_paragraph(clean(line.lstrip("> ")), style="Subtitle")
            p.paragraph_format.space_after = Pt(5)
            continue
        if re.match(r"^\s*[-*]\s+", line):
            body = re.sub(r"^\s*[-*]\s+", "", line)
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(clean(body))
            continue
        numbered = re.match(r"^\s*(\d+)[.)]\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.add_run(f"{numbered.group(1)}.  {clean(numbered.group(2))}")
            continue
        p = doc.add_paragraph(clean(line))
        if line.startswith("**") or line.startswith("- **"):
            p.runs[0].bold = True


def add_bundle(doc: Document) -> None:
    main = (SRS / "OvO-SRS-Draft.md").read_text(encoding="utf-8")
    sections = re.split(r"(?=^## \d+ )", main, flags=re.M)
    preamble = sections.pop(0)
    # The cover carries the title and release; retain introductory source notes.
    temp_dir = ROOT / ".local" / "docx-build"
    temp_dir.mkdir(parents=True, exist_ok=True)
    preamble_file = temp_dir / "preamble.md"
    preamble_file.write_text(preamble, encoding="utf-8")
    add_markdown(doc, preamble_file, skip_title=True)
    inserts = {
        3: [("用例独立规约（UC-01～UC-15）", "use-case-specifications.md", 1)],
        4: [
            ("功能详细规格：检索与门户", "modules/detailed/01-search-and-portal-detailed.md", 1),
            ("功能详细规格：账号、认领与维护", "modules/detailed/02-account-claim-maintenance-detailed.md", 1),
            ("功能详细规格：采集、治理与更新", "modules/detailed/03-collection-governance-detailed.md", 1),
            ("功能详细规格：统计、管理与创新", "modules/detailed/04-analytics-admin-innovation-detailed.md", 1),
        ],
        6: [("权限、状态与规则适用矩阵", "cross-cutting-matrices.md", 1)],
        7: [
            ("业务规则 BR-01～BR-20", "business-rules.md", 1),
            ("数据定义与细化规则", "data-and-rules.md", 1),
        ],
        9: [("非功能逐条规格 NF01～NF08", "modules/detailed/nonfunctional-detailed.md", 1)],
        10: [
            ("父需求追踪矩阵", "traceability-matrix.md", 1),
            ("124条子需求追踪矩阵", "subrequirement-traceability.md", 1),
        ],
    }
    for section in sections:
        number = int(re.match(r"## (\d+)", section).group(1))
        file = temp_dir / f"section-{number}.md"
        file.write_text(section, encoding="utf-8")
        add_markdown(doc, file, base=-1)
        for label, relative, base in inserts.get(number, []):
            heading = doc.add_paragraph(label, style="Heading 2")
            heading.paragraph_format.keep_with_next = True
            add_markdown(doc, SRS / relative, base=base, skip_title=True)


def build() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.1), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.35), Cm(2.1)
    sec.header_distance, sec.footer_distance = Cm(1), Cm(0.9)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.space_after = Pt(4)
    for idx, size in ((1, 16), (2, 13), (3, 10.5)):
        style = styles[f"Heading {idx}"]
        style.font.name = "微软雅黑"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.paragraph_format.space_before = Pt(14 if idx == 1 else 8)
        style.paragraph_format.space_after = Pt(6 if idx == 1 else 4)
    for name in ("List Bullet", "List Number"):
        styles[name].font.name = "宋体"
        styles[name].font.size = Pt(10)
        styles[name].paragraph_format.space_after = Pt(2)
    title_style = styles["Title"]
    title_style.font.name = "微软雅黑"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title_style.font.size = Pt(22)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title_style._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(120)
    r = title.add_run("OvO 学术成果分享平台\n软件需求规格说明书")
    r.bold = True
    r.font.name = "微软雅黑"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0, 0, 0)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(40)
    meta.add_run("V0.4 · 提交评审稿\n2026 年 10 月 12 日\nOvO 小组").font.size = Pt(12)
    doc.add_paragraph("组号、成员分工、成员复核与共同冻结：待团队填写或签认。两张状态图已有图源，其他图与CodeArts一致性仍待核验。")
    doc.add_page_break()
    p = doc.add_paragraph("文档结构", style="Heading 1")
    p.paragraph_format.space_after = Pt(10)
    for text in (
        "1 引言", "2 系统总体说明", "3 用例模型", "4 功能需求规格", "5 创新与可选需求",
        "6 业务流程与状态模型", "7 数据需求与业务规则", "8 权限与安全需求", "9 非功能需求", "10 验收与追踪",
    ):
        doc.add_paragraph(text)
    doc.add_paragraph("各章详细子需求、用例与追踪矩阵紧随所属章节。可使用 Word 导航窗格按标题定位。")
    doc.add_page_break()
    add_bundle(doc)
    header = sec.header.paragraphs[0]
    header.text = "OvO 学术成果分享平台  |  SRS V0.4"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.color.rgb = RGBColor(95, 110, 125)
    header.runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("提交评审稿 · 第 ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    footer.add_run(" 页")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
